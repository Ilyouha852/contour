from __future__ import annotations

import csv
import html
import io
import json
from typing import Any
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .store import connection


def _rows_for_run(run_id: str) -> list[dict[str, Any]]:
    with connection() as conn:
        rows = conn.execute(
            """SELECT lot_id,rank,supplier_inn,supplier_name,score,role,role_conf,role_source,role_signals_json,status,is_msp,risk,
                      msp_status,risk_status,
                      enrichment_json,factors_json,factor_scores_json,evidence_json,explanation
            FROM recommendations WHERE run_id=? ORDER BY lot_id,rank""",
            (run_id,),
        ).fetchall()
    return [dict(row) for row in rows]


def export_csv(run_id: str) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=";")
    writer.writerow(
        [
            "lot_id",
            "rank",
            "supplier_inn",
            "name",
            "score",
            "role",
            "role_conf",
            "role_source",
            "role_signals",
            "status",
            "msp",
            "risk",
                "msp_status",
                "risk_status",
                "enrichment",
            "factors",
            "factor_scores",
            "evidence",
            "explanation",
        ]
    )
    for row in _rows_for_run(run_id):
        writer.writerow(
            [
                row["lot_id"],
                row["rank"],
                row["supplier_inn"],
                row["supplier_name"] or "",
                row["score"],
                row["role"],
                row["role_conf"],
                row["role_source"],
                row["role_signals_json"],
                row["status"],
                    "true" if row["is_msp"] else "false" if row["msp_status"] == "not_member" else "unknown",
                    "true" if row["risk"] else "false" if row["risk_status"] == "clear" else "unknown",
                    row["msp_status"],
                    row["risk_status"],
                row["enrichment_json"],
                row["factors_json"],
                row["factor_scores_json"],
                row["evidence_json"],
                row["explanation"],
            ]
        )
    return b"\xef\xbb\xbf" + output.getvalue().encode("utf-8")


def export_xlsx(run_id: str) -> bytes:
    rows = _rows_for_run(run_id)
    workbook = Workbook()
    recommendations = workbook.active
    recommendations.title = "Рекомендации"
    recommendations.append(
            ["Лот", "Ранг", "ИНН", "Наименование", "Оценка", "Роль", "Уверенность", "Статус", "МСП", "Риск", "Проверка МСП", "Проверка РНП"]
    )
    explanations = workbook.create_sheet("Объяснения")
    explanations.append(
        [
            "Лот",
            "Ранг",
            "ИНН",
            "Источник классификации",
            "Сигналы роли",
            "Источники и дата проверки",
            "Факторы (вклад)",
            "Оценки факторов",
            "Лоты-доказательства",
            "Объяснение",
        ]
    )

    for row in rows:
        recommendations.append(
            [
                row["lot_id"],
                row["rank"],
                row["supplier_inn"],
                row["supplier_name"],
                row["score"],
                row["role"],
                row["role_conf"],
                row["status"],
                    "true" if row["is_msp"] else "false" if row["msp_status"] == "not_member" else "unknown",
                    "true" if row["risk"] else "false" if row["risk_status"] == "clear" else "unknown",
                    row["msp_status"],
                    row["risk_status"],
            ]
        )
        explanations.append(
            [
                row["lot_id"],
                row["rank"],
                row["supplier_inn"],
                row["role_source"],
                row["role_signals_json"],
                row["enrichment_json"],
                row["factors_json"],
                row["factor_scores_json"],
                row["evidence_json"],
                row["explanation"],
            ]
        )

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243B53")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column_cells in sheet.columns:
            letter = get_column_letter(column_cells[0].column)
            sample = [str(cell.value or "") for cell in list(column_cells)[:100]]
            width = min(60, max(12, max((len(value) for value in sample), default=12) + 2))
            sheet.column_dimensions[letter].width = width
        sheet.row_dimensions[1].height = 30

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_passport_html(run_id: str) -> str | None:
    with connection() as conn:
        run = conn.execute("SELECT * FROM runs WHERE run_id=?", (run_id,)).fetchone()
        lots = {
            row["lot_id"]: dict(row)
            for row in conn.execute(
                """SELECT DISTINCT l.lot_id,l.publish_date,l.start_price,l.procedure_name,l.subject,l.customer_inn
                FROM lots l JOIN recommendations r ON r.lot_id=l.lot_id
                WHERE r.run_id=?""",
                (run_id,),
            )
        }
    if run is None:
        return None

    rows = _rows_for_run(run_id)
    weights = json.loads(run["weights_json"])
    factor_names = {
        "F1": "Опыт по позициям",
        "F2": "Текстовое сходство",
        "F3": "История побед",
        "F4": "Опыт у заказчика",
        "F5": "Ценовой диапазон",
        "F6": "Давность победы",
        "F7": "Регион",
        "F8": "МСП",
    }
    formulas = {
        "F1": "1 − exp(−взвешенный опыт / τ)",
        "F2": "min(1, текстовые совпадения / вес выборки)",
        "F3": "(похожие победы + α × базовая доля побед) / (похожие участия + α)",
        "F4": "1 − exp(−победы у заказчика / 2)",
        "F5": "exp(−|ln(НМЦК) − ln(медиана цен)| / σ)",
        "F6": "exp(−дней с похожей победы / 730)",
        "F7": "Совпадение регионов поставщика и заказчика",
        "F8": "Для закупки МСП: member=1, not_member=0, unknown=0.5",
    }

    def escape(value: Any) -> str:
        return html.escape(str(value if value is not None else "—"), quote=True)

    def money(value: Any) -> str:
        if value is None:
            return "Не указана"
        return f"{float(value):,.0f} ₽".replace(",", " ")

    def safe_href(value: Any) -> str | None:
        if not isinstance(value, str):
            return None
        parsed = urlsplit(value)
        return value if parsed.scheme in {"http", "https"} and parsed.netloc else None

    def render_source(title: str, fact: dict[str, Any]) -> str:
        source_url = safe_href(fact.get("source_url"))
        source = escape(fact.get("source"))
        link = f'<a href="{escape(source_url)}">Открыть источник</a>' if source_url else ""
        return (
            f'<li><strong>{escape(title)}: {escape(fact.get("status"))}</strong> · {source} · '
            f'проверено {escape(fact.get("checked_at"))} {link}</li>'
        )

    cards = []
    for lot_id, lot_rows in _group_rows_by_lot(rows).items():
        lot = lots.get(lot_id, {})
        title = lot.get("procedure_name") or lot.get("subject") or f"Лот {lot_id}"
        candidate_cards = []
        for row in lot_rows[:10]:
            scores = json.loads(row["factor_scores_json"])
            contributions = json.loads(row["factors_json"])
            evidence = json.loads(row["evidence_json"])
            enrichment = json.loads(row["enrichment_json"])
            total_contribution = sum(float(value) for value in contributions.values())
            multiplier = min(1.0, max(0.0, float(row["score"]) / (100.0 * total_contribution))) if total_contribution else 0.0
            factor_rows = "".join(
                "<tr>"
                f"<td>{escape(factor_names.get(factor, factor))}</td>"
                f"<td>{escape('100%, так как лот не ограничен для МСП' if factor == 'F8' and lot.get('is_smp') != 1 else formulas.get(factor, ''))}</td>"
                f"<td>{float(scores.get(factor, 0)):.0%}</td>"
                f"<td>{float(weights.get(factor, 0)):.0%}</td>"
                f"<td>{100.0 * float(contributions.get(factor, 0)) * multiplier:.2f}</td>"
                "</tr>"
                for factor in sorted(scores)
            )
            evidence_rows = "".join(
                "<li>"
                f"<strong>Лот {escape(item.get('lot_id'))}</strong> · {escape(item.get('publish_date'))} · "
                f"{escape(item.get('title'))} · ОКПД2 {escape(', '.join(item.get('okpd2') or []))} · "
                f"покрытие {float(item.get('item_coverage') or 0):.0%} · "
                f"текст {float(item.get('text_similarity') or 0):.0%} · "
                f"цена {escape(money(item.get('price')))}"
                "</li>"
                for item in evidence
            ) or "<li>Подтверждающие победы не найдены.</li>"
            sources = "".join(
                render_source(label, enrichment.get(key) or {})
                for key, label in (("msp", "МСП"), ("rnp", "РНП"))
            )
            candidate_cards.append(
                f"""<article class="candidate">
<div class="candidate-head"><span class="rank">#{escape(row['rank'])}</span><div><h3>{escape(row['supplier_name'] or 'Поставщик')}</h3><p>ИНН {escape(row['supplier_inn'])} · {escape(row['role'])} · {escape(row['status'])}</p></div><strong class="score">{float(row['score']):.1f}</strong></div>
<p class="explanation">{escape(row['explanation'])}</p>
<h4>Факторы и формулы</h4><table><thead><tr><th>Фактор</th><th>Формула</th><th>Значение</th><th>Вес</th><th>Вклад, баллы</th></tr></thead><tbody>{factor_rows}</tbody></table>
<h4>Лоты-доказательства</h4><ul>{evidence_rows}</ul>
<h4>Реестры и provenance</h4><ul>{sources}</ul>
</article>"""
            )
        cards.append(
            f"""<section class="lot">
<header><p class="eyebrow">ЗАКУПКА · {escape(lot.get('publish_date'))}</p><h2>Лот {escape(lot_id)} · {escape(title)}</h2><p>НМЦК {escape(money(lot.get('start_price')))} · заказчик ИНН {escape(lot.get('customer_inn'))}</p></header>
{''.join(candidate_cards)}
</section>"""
        )

    return f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Паспорт доказательств · {escape(run_id)}</title>
<style>
*{{box-sizing:border-box}}body{{margin:0;background:#f3f7f5;color:#203330;font:14px/1.5 "Segoe UI",sans-serif}}main{{width:min(1040px,calc(100% - 40px));margin:32px auto}}.toolbar{{display:flex;justify-content:flex-end;margin-bottom:14px}}button{{padding:9px 13px;border:1px solid #b9d0c7;background:#087668;color:#fff;font:inherit;font-weight:700;cursor:pointer}}header,.candidate{{padding:18px 20px;background:#fff;border:1px solid #dce6e2;margin-bottom:12px}}h1,h2,h3,h4,p{{margin-top:0}}h1{{font:500 30px Georgia,serif}}h2{{font-size:18px;line-height:1.35}}h3{{margin:0;font-size:15px}}h4{{margin:16px 0 6px;font-size:12px}}.eyebrow{{color:#087668;font-size:10px;font-weight:800;letter-spacing:.08em}}.meta,.explanation,.candidate-head p{{color:#657773;font-size:11px}}.candidate-head{{display:flex;align-items:center;gap:12px}}.rank{{display:grid;place-items:center;width:32px;height:32px;background:#e4f2ee;color:#07594f;font-weight:800}}.candidate-head>div{{min-width:0;flex:1}}.score{{color:#176b5d;font-size:20px}}table{{width:100%;border-collapse:collapse;font-size:10px}}th,td{{padding:6px;text-align:left;vertical-align:top;border-bottom:1px solid #e5ece8}}th{{background:#f5f8f6;color:#52655d}}ul{{padding-left:18px;margin:4px 0;font-size:10px}}li{{margin:3px 0;overflow-wrap:anywhere}}a{{color:#087668}}.note{{margin-top:20px;color:#71817b;font-size:10px}}@media print{{body{{background:#fff}}main{{width:100%;margin:0}}.toolbar{{display:none}}header,.candidate{{break-inside:avoid;border-color:#cbd5d1;margin-bottom:8px;padding:12px}}h1{{font-size:24px}}}}
</style></head><body><main>
<div class="toolbar"><button onclick="window.print()">Печать / сохранить PDF</button></div>
<header><p class="eyebrow">КОНТУР ЗАКУПОК · ПАСПОРТ ДОКАЗАТЕЛЬСТВ</p><h1>Обоснование рекомендаций</h1><p>{escape(run["created_at"])} · данные {escape(run["data_version"])} · запуск {escape(run_id)}</p><p class="meta">Формула итоговой оценки: 100 × сумма(значение фактора × вес) × коэффициент истории. Неизвестный статус реестра не трактуется как отсутствие записи.</p></header>
{''.join(cards)}
<p class="note">Источники реестров приведены вместе со временем проверки. Паспорт отражает сохранённый результат запуска.</p>
</main></body></html>"""


def _group_rows_by_lot(rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        grouped.setdefault(row["lot_id"], []).append(row)
    return grouped


def parse_run_row(row) -> dict[str, Any]:
    result = dict(row)
    result["weights"] = json.loads(result.pop("weights_json"))
    total = result["n_lots"]
    result["progress_percent"] = round(100 * result["completed_lots"] / total, 1) if total else 100.0
    return result